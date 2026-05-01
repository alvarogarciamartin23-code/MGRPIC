from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Márgenes A4 estándar TFM ──────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Cm(21)
section.page_height = Cm(29.7)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(2.0)

# ── Paleta de colores ─────────────────────────────────────────────────────────
AZUL       = RGBColor(0x1F, 0x38, 0x64)   # institucional
AZUL_CLARO = RGBColor(0x2E, 0x74, 0xB5)   # acento
MORADO     = RGBColor(0x6C, 0x34, 0x83)   # D1
AZUL_D2    = RGBColor(0x1A, 0x52, 0x76)   # D2
ROJO_D3    = RGBColor(0x92, 0x2B, 0x21)   # D3
VERDE      = RGBColor(0x70, 0xAD, 0x47)
AMBAR      = RGBColor(0xFF, 0xD9, 0x66)
NARANJA    = RGBColor(0xFF, 0x99, 0x33)
ROJO       = RGBColor(0xFF, 0x00, 0x00)
GRIS_FONDO = RGBColor(0xF2, 0xF2, 0xF2)
NEGRO      = RGBColor(0x1C, 0x25, 0x35)
GRIS_TEXT  = RGBColor(0x4A, 0x5E, 0x72)

def set_cell_bg(cell, rgb_hex):
    """Establece color de fondo de celda."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), rgb_hex)
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if val:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), val.get('val', 'single'))
            el.set(qn('w:sz'), str(val.get('sz', 4)))
            el.set(qn('w:color'), val.get('color', 'auto'))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def par_style(par, size=11, bold=False, italic=False, color=None, align=None, space_before=0, space_after=6):
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after  = Pt(space_after)
    if align: par.alignment = align
    for run in par.runs:
        run.font.size   = Pt(size)
        run.font.bold   = bold
        run.font.italic = italic
        if color: run.font.color.rgb = color

def add_par(doc, text, size=11, bold=False, italic=False, color=None,
            align=WD_ALIGN_PARAGRAPH.LEFT, sb=0, sa=4, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    p.alignment = align
    run = p.add_run(text)
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.name   = 'Calibri'
    if color: run.font.color.rgb = color
    return p

def add_heading_annex(doc, text, level=1):
    """Encabezado numerado para secciones del anexo."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level==1 else 8)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.font.name  = 'Calibri'
    run.font.size  = Pt(13 if level==1 else 11)
    run.font.bold  = True
    run.font.color.rgb = AZUL if level==1 else AZUL_CLARO
    # Línea inferior para nivel 1
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:color'), '1F3864')
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


# ══════════════════════════════════════════════════════════════════════════════
# PORTADA DEL ANEXO
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()

p = add_par(doc, 'ANEXO', size=11, bold=False, color=GRIS_TEXT,
            align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=4)
p.runs[0].font.letter_spacing = Pt(2)

add_par(doc, 'MGRPIC — Herramienta de Evaluación de Riesgo Procesal',
        size=18, bold=True, color=AZUL, align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=6)

add_par(doc, 'Matriz de Gestión de Riesgo Procesal en la Incautación de Criptoactivos',
        size=13, bold=False, color=AZUL_CLARO, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=4)

add_par(doc, 'Versión 3.0 · Instrumento de uso policial',
        size=10, italic=True, color=GRIS_TEXT, align=WD_ALIGN_PARAGRAPH.CENTER, sb=8, sa=4)

doc.add_paragraph()

# Nota de archivo digital
t = doc.add_table(rows=1, cols=1)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.style = 'Table Grid'
cell = t.cell(0, 0)
set_cell_bg(cell, 'E6EFF8')
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after  = Pt(6)
r = p.add_run('El archivo ejecutable ')
r.font.size = Pt(10); r.font.name = 'Calibri'
r2 = p.add_run('MGRPIC.html')
r2.font.size = Pt(10); r2.font.bold = True; r2.font.name = 'Calibri'
r3 = p.add_run(' se adjunta en formato digital. Abrir con cualquier navegador web sin instalación ni conexión a internet.')
r3.font.size = Pt(10); r3.font.name = 'Calibri'

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 1. DESCRIPCIÓN TÉCNICA
# ══════════════════════════════════════════════════════════════════════════════
add_heading_annex(doc, '1. Descripción técnica de la herramienta')

rows_desc = [
    ('Denominación',        'Matriz de Gestión de Riesgo Procesal en la Incautación de Criptoactivos (MGRPIC)'),
    ('Versión',             '3.0'),
    ('Tipo de aplicación',  'Aplicación web de página única (SPA). Archivo HTML autocontenido sin dependencias externas.'),
    ('Tecnología',          'HTML5, CSS3 y JavaScript puro (sin frameworks ni librerías externas). Funciona completamente offline.'),
    ('Compatibilidad',      'Cualquier navegador web moderno: Chrome, Firefox, Edge, Safari (versiones actuales).'),
    ('Instalación',         'No requiere instalación. Basta con abrir el archivo MGRPIC.html en el navegador.'),
    ('Almacenamiento',      'Ningún dato se almacena en el dispositivo ni se transmite a servidores externos. El estado de la evaluación se mantiene únicamente en memoria durante la sesión.'),
    ('Naturaleza',          'Instrumento de uso policial. Apoya dos decisiones: selección del modelo de custodia y valoración de la procedencia de propuesta de enajenación anticipada al Ministerio Fiscal al amparo del artículo 367 ter de la LECrim. No es una herramienta judicial.'),
]

t = doc.add_table(rows=len(rows_desc), cols=2)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.LEFT
col_widths = [Cm(4.5), Cm(11.0)]

for i, (lbl, val) in enumerate(rows_desc):
    row = t.rows[i]
    row.cells[0].width = col_widths[0]
    row.cells[1].width = col_widths[1]
    set_cell_bg(row.cells[0], 'E6EFF8')
    # Label
    p0 = row.cells[0].paragraphs[0]
    p0.paragraph_format.space_before = Pt(3)
    p0.paragraph_format.space_after  = Pt(3)
    r = p0.add_run(lbl)
    r.font.size = Pt(9.5); r.font.bold = True; r.font.name = 'Calibri'
    r.font.color.rgb = AZUL
    # Value
    p1 = row.cells[1].paragraphs[0]
    p1.paragraph_format.space_before = Pt(3)
    p1.paragraph_format.space_after  = Pt(3)
    r2 = p1.add_run(val)
    r2.font.size = Pt(9.5); r2.font.name = 'Calibri'

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 2. ESTRUCTURA DE INDICADORES
# ══════════════════════════════════════════════════════════════════════════════
add_heading_annex(doc, '2. Estructura de indicadores y sistema de ponderación')

add_par(doc, 'La herramienta evalúa 10 indicadores agrupados en 3 dimensiones ponderadas. '
        'Cada indicador se puntúa en una escala de 0 (Bajo), 1 (Moderado) o 2 (Alto). '
        'La puntuación global se expresa en una escala de 0 a 100 puntos.',
        size=10, color=NEGRO, sb=0, sa=8)

# Tabla de dimensiones e indicadores
dims = [
    {
        'id': 'D1', 'nombre': 'Volatilidad', 'peso': '35 %', 'factor': '4,375',
        'max': '35', 'color': '6C3483', 'color_rgb': MORADO,
        'indicadores': [
            ('V.1', 'Naturaleza del activo',
             'Stablecoin (USDT, USDC)',
             'Alta capitalización con mercado líquido (BTC, ETH)',
             'Altcoin de baja capitalización o token sin mercado secundario activo'),
            ('V.2', 'Variación de precio en los 30 días previos a la incautación',
             'Inferior al 10% sin tendencia marcada',
             'Entre el 10% y el 30%',
             'Superior al 30% o variación brusca >15% en las últimas 48 horas'),
            ('V.3', 'Complejidad de la causa como indicador de duración estimada',
             'Causa sencilla, investigado único, sin cooperación internacional',
             'Varios investigados o diligencia internacional pendiente',
             'Causa compleja, cooperación internacional activa o crimen organizado'),
            ('V.4', 'Valor total del activo incautado',
             'Inferior a 10.000 €',
             'Entre 10.000 € y 500.000 €',
             'Superior a 500.000 €'),
        ]
    },
    {
        'id': 'D2', 'nombre': 'Custodia y Trazabilidad', 'peso': '30 %', 'factor': '5,0',
        'max': '30', 'color': '1A5276', 'color_rgb': AZUL_D2,
        'indicadores': [
            ('C.1', 'Tipo de soporte de custodia de la clave privada',
             'Custodio institucional especializado (Prosegur Crypto o ORGA) con contrato',
             'Monedero frío policial bajo acta formal de entrega al LAJ',
             'Soporte en papel, dispositivo sin cifrar o custodia informal sin acta'),
            ('C.2', 'Número de personas con conocimiento de la clave privada',
             'Una sola persona con acta de entrega y registro nominativo',
             'Dos o tres personas con registro documentado de acceso',
             'Más de tres personas o acceso sin documentación formal'),
            ('C.3', 'Calidad de la trazabilidad del activo en la cadena de bloques',
             'Dirección vinculada directamente al investigado mediante datos del exchange con KYC',
             'Dirección identificada mediante heurísticas de clustering, sin vinculación directa',
             'Activo trazado parcialmente o flujo interrumpido por mixer, tumbler o moneda de privacidad'),
        ]
    },
    {
        'id': 'D3', 'nombre': 'Contexto Operativo Policial', 'peso': '35 %', 'factor': '5,833',
        'max': '35', 'color': '922B21', 'color_rgb': ROJO_D3,
        'indicadores': [
            ('J.1', 'Existencia de protocolo previo en la unidad para intervención de criptoactivos',
             'Unidad con protocolo documentado y personal con formación acreditada',
             'Sin protocolo propio pero con acceso a unidad especializada de apoyo',
             'Sin protocolo, sin personal formado y sin unidad especializada de apoyo'),
            ('J.2', 'Urgencia de la decisión de enajenación anticipada',
             'Activo estable y causa corta; custodia ordinaria sin riesgo patrimonial relevante',
             'Volatilidad moderada o causa larga; conveniente elevar propuesta al Fiscal en plazo breve',
             'Alta volatilidad o valor elevado con causa larga; riesgo patrimonial inmediato'),
            ('J.3', 'Claridad de la cadena de responsabilidad sobre el activo incautado',
             'Responsabilidad claramente atribuida al custodio institucional mediante contrato formal',
             'Responsabilidad compartida entre LAJ y Policía Judicial con acta de entrega documentada',
             'Custodia sin atribución clara de responsabilidad o sin documentación formal'),
        ]
    },
]

for dim in dims:
    # Encabezado de dimensión
    n_ind = len(dim['indicadores'])
    # Header row: Dim ID | Nombre | Peso | Factor | Máx.
    t = doc.add_table(rows=1 + n_ind, cols=6)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Anchos de columna
    widths = [Cm(0.9), Cm(1.1), Cm(3.8), Cm(3.0), Cm(3.0), Cm(3.0)]
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]

    # Fila de encabezado de la dimensión (fila 0)
    hrow = t.rows[0]
    set_cell_bg(hrow.cells[0], dim['color'])
    set_cell_bg(hrow.cells[1], dim['color'])
    set_cell_bg(hrow.cells[2], dim['color'])
    set_cell_bg(hrow.cells[3], 'E6EFF8')
    set_cell_bg(hrow.cells[4], 'E6EFF8')
    set_cell_bg(hrow.cells[5], 'E6EFF8')

    # Merge cells 0+1+2 para nombre de dimensión
    hrow.cells[0].merge(hrow.cells[2])
    p = hrow.cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(f'{dim["id"]} — {dim["nombre"]}')
    r.font.size = Pt(10); r.font.bold = True; r.font.name = 'Calibri'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for col_idx, (label, val) in enumerate([
        ('Peso', dim['peso']), ('Factor', dim['factor']), ('Máx.', dim['max'] + ' pts')
    ], start=3):
        p = hrow.cells[col_idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(f'{label}: {val}')
        r.font.size = Pt(9); r.font.bold = True; r.font.name = 'Calibri'
        r.font.color.rgb = AZUL

    # Sub-encabezado de columnas de indicadores
    # Fila 1..n: indicadores
    col_headers = ['Cód.', 'Nombre del indicador', 'Bajo (0)', 'Moderado (1)', 'Alto (2)']
    col_bg = ['E6EFF8', 'E6EFF8', 'E8F5DD', 'FFF8CC', 'FFE5E5']

    for i_ind, (codigo, nombre, bajo, mod, alto) in enumerate(dim['indicadores']):
        row = t.rows[1 + i_ind]
        vals = [codigo, nombre, bajo, mod, alto]
        bgs  = col_bg
        for ci, (v, bg) in enumerate(zip(vals, bgs)):
            set_cell_bg(row.cells[ci + 1], bg)
            p = row.cells[ci + 1].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(3)
            if ci == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(v)
            r.font.size = Pt(8.5); r.font.name = 'Calibri'
            r.font.bold = (ci <= 1)
            if ci == 0: r.font.color.rgb = dim['color_rgb']
        # Celda 0 (vacía, borde de dimensión)
        set_cell_bg(row.cells[0], dim['color'])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ══════════════════════════════════════════════════════════════════════════════
# 3. NIVELES DE RIESGO Y PROTOCOLO POLICIAL
# ══════════════════════════════════════════════════════════════════════════════
add_heading_annex(doc, '3. Niveles de riesgo y protocolo policial de actuación')

niveles = [
    ('BAJO',     '0 – 25',   'E8F5DD', '70AD47', '3D7020',
     'Riesgo procesal controlado. Custodia ordinaria bajo supervisión del LAJ. Incorporar resultado '
     'al atestado. Reevaluar si algún indicador cambia significativamente, en particular ante '
     'variaciones relevantes del precio del activo.'),
    ('MODERADO', '26 – 50',  'FFF8CC', 'B8860B', '7D5C00',
     'Riesgo procesal significativo. Valorar encomienda a la ORGA o custodio institucional especializado. '
     'Documentar variaciones de precio mensualmente. Si V.2 asciende a Alto, elevar propuesta motivada '
     'al Ministerio Fiscal para que inste la enajenación anticipada conforme al artículo 367 ter LECrim.'),
    ('ALTO',     '51 – 75',  'FFF0D9', 'FF9933', 'A04800',
     'Riesgo procesal elevado. Elevar propuesta motivada al Ministerio Fiscal en el plazo más breve '
     'posible, con fundamento en la evaluación MGRPIC. Solicitar perito especializado en activos '
     'digitales. Valorar la enajenación anticipada o conversión a moneda fiduciaria estable.'),
    ('CRÍTICO',  '76 – 100', 'FFE5E5', 'FF0000', 'CC0000',
     'Riesgo procesal máximo. Actuación inmediata en las primeras 24-48 horas. Comunicación urgente '
     'al Ministerio Fiscal. Elevación inmediata de propuesta de enajenación anticipada al amparo del '
     'artículo 367 ter LECrim. Apertura de pieza separada de responsabilidad civil. Valorar el bloqueo '
     'cautelar de activos en exchanges identificados.'),
]

t = doc.add_table(rows=len(niveles) + 1, cols=3)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.LEFT
widths_n = [Cm(2.2), Cm(2.0), Cm(11.3)]

# Encabezado
for ci, (hdr, bg) in enumerate(zip(['Nivel', 'Puntuación', 'Protocolo policial de actuación'],
                                    ['1F3864', '1F3864', '1F3864'])):
    set_cell_bg(t.rows[0].cells[ci], bg)
    t.rows[0].cells[ci].width = widths_n[ci]
    p = t.rows[0].cells[ci].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(hdr)
    r.font.size = Pt(9); r.font.bold = True; r.font.name = 'Calibri'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for i, (nivel, rango, bg, border_color, texto_color, protocolo) in enumerate(niveles):
    row = t.rows[i + 1]
    for ci in range(3): row.cells[ci].width = widths_n[ci]
    set_cell_bg(row.cells[0], bg)
    set_cell_bg(row.cells[1], bg)
    set_cell_bg(row.cells[2], 'FFFFFF')

    p0 = row.cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after  = Pt(4)
    r = p0.add_run(nivel)
    r.font.size = Pt(9.5); r.font.bold = True; r.font.name = 'Calibri'
    tc = int(texto_color[0:2],16), int(texto_color[2:4],16), int(texto_color[4:6],16)
    r.font.color.rgb = RGBColor(*tc)

    p1 = row.cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after  = Pt(4)
    r2 = p1.add_run(rango)
    r2.font.size = Pt(9.5); r2.font.bold = True; r2.font.name = 'Calibri'
    r2.font.color.rgb = RGBColor(*tc)

    p2 = row.cells[2].paragraphs[0]
    p2.paragraph_format.space_before = Pt(3)
    p2.paragraph_format.space_after  = Pt(3)
    r3 = p2.add_run(protocolo)
    r3.font.size = Pt(9); r3.font.name = 'Calibri'

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 4. VERIFICACIÓN MATEMÁTICA — CASOS TIPO
# ══════════════════════════════════════════════════════════════════════════════
add_heading_annex(doc, '4. Verificación matemática del modelo de ponderación')

add_par(doc, 'La validez del sistema de ponderación se ha verificado mediante cinco casos tipo que cubren '
        'los extremos y los valores intermedios de la escala. Los resultados obtenidos por la herramienta '
        'coinciden con los valores calculados manualmente.',
        size=10, color=NEGRO, sb=0, sa=8)

casos = [
    ('A', 'V1=1, V2=0, V3=1, V4=1', 'C1=0, C2=0, C3=0', 'J1=1, J2=0, J3=0',
     '13,13', '0,00', '5,83', '18,96', 'BAJO'),
    ('B', 'V1=2, V2=2, V3=2, V4=2', 'C1=2, C2=2, C3=2', 'J1=2, J2=2, J3=2',
     '35,00', '30,00', '35,00', '100,00', 'CRÍTICO'),
    ('C', 'V1=1, V2=1, V3=2, V4=2', 'C1=1, C2=1, C3=1', 'J1=1, J2=1, J3=1',
     '26,25', '15,00', '17,50', '58,75', 'ALTO'),
    ('D', 'V1=0, V2=0, V3=0, V4=1', 'C1=0, C2=0, C3=0', 'J1=1, J2=0, J3=0',
     '4,38', '0,00', '5,83', '10,21', 'BAJO'),
    ('E', 'V1=1, V2=2, V3=2, V4=2', 'C1=1, C2=2, C3=2', 'J1=2, J2=2, J3=1',
     '30,63', '25,00', '29,17', '84,79', 'CRÍTICO'),
]

hdrs_c = ['Caso', 'D1 — Volatilidad', 'D2 — Custodia y traz.', 'D3 — Contexto op.', 'Pts. D1', 'Pts. D2', 'Pts. D3', 'Total', 'Nivel']
bgs_c  = ['1F3864','6C3483','1A5276','922B21','E6EFF8','E6EFF8','E6EFF8','E6EFF8','1F3864']

t = doc.add_table(rows=len(casos)+1, cols=len(hdrs_c))
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.LEFT
w_c = [Cm(0.8), Cm(2.6), Cm(2.6), Cm(2.6), Cm(1.1), Cm(1.1), Cm(1.1), Cm(1.1), Cm(1.5)]

for ci, (hdr, bg) in enumerate(zip(hdrs_c, bgs_c)):
    cell = t.rows[0].cells[ci]
    cell.width = w_c[ci]
    set_cell_bg(cell, bg)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(hdr)
    r.font.size = Pt(8); r.font.bold = True; r.font.name = 'Calibri'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if bg != 'E6EFF8' else AZUL

nivel_colors = {
    'BAJO': ('E8F5DD', '3D7020'), 'MODERADO': ('FFF8CC', '7D5C00'),
    'ALTO': ('FFF0D9', 'A04800'), 'CRÍTICO': ('FFE5E5', 'CC0000')
}

for i, (caso, d1, d2, d3, p1, p2, p3, total, nivel) in enumerate(casos):
    row = t.rows[i+1]
    for ci in range(len(hdrs_c)): row.cells[ci].width = w_c[ci]
    vals = [caso, d1, d2, d3, p1, p2, p3, total, nivel]
    bg_c, txt_c = nivel_colors[nivel]
    for ci, v in enumerate(vals):
        cell = row.cells[ci]
        if ci == 0:
            set_cell_bg(cell, bg_c)
        elif ci == 8:
            set_cell_bg(cell, bg_c)
        else:
            set_cell_bg(cell, 'FFFFFF')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(v)
        r.font.size = Pt(8); r.font.name = 'Calibri'
        r.font.bold = (ci in [0, 7, 8])
        if ci in [0, 8]:
            tc2 = int(txt_c[0:2],16), int(txt_c[2:4],16), int(txt_c[4:6],16)
            r.font.color.rgb = RGBColor(*tc2)

doc.add_paragraph()

add_par(doc,
    'Fórmulas aplicadas: D1 = suma(V1–V4) × 4,375 | D2 = suma(C1–C3) × 5,000 | '
    'D3 = suma(J1–J3) × 5,833 | Puntuación global = D1 + D2 + D3 (escala 0–100)',
    size=8.5, italic=True, color=GRIS_TEXT, sb=2, sa=8)

# ══════════════════════════════════════════════════════════════════════════════
# 5. INSTRUCCIONES DE USO
# ══════════════════════════════════════════════════════════════════════════════
add_heading_annex(doc, '5. Instrucciones de uso')

pasos = [
    ('Paso 1 — Datos del caso',
     'Introducir el número de diligencias, la unidad policial actuante, la fecha de la incautación, '
     'el tipo de activo digital intervenido y el órgano judicial competente. Estos datos se integrarán '
     'automáticamente en el encabezado del informe final.'),
    ('Paso 2 — Dimensión I: Volatilidad',
     'Seleccionar una opción (Bajo / Moderado / Alto) para cada uno de los cuatro indicadores V.1 a V.4. '
     'Es obligatorio responder a todos los indicadores antes de continuar. '
     'Un contador indica el progreso de la dimensión.'),
    ('Paso 3 — Dimensión II: Custodia y Trazabilidad',
     'Seleccionar una opción para cada uno de los tres indicadores C.1 a C.3, relativos al modelo '
     'de custodia de la clave privada y a la calidad de la trazabilidad del activo en blockchain.'),
    ('Paso 4 — Dimensión III: Contexto Operativo Policial',
     'Seleccionar una opción para cada uno de los tres indicadores J.1 a J.3, relativos a la capacidad '
     'operativa de la unidad, la urgencia de la decisión de enajenación y la cadena de responsabilidad.'),
    ('Paso 5 — Resultados',
     'La herramienta calcula automáticamente la puntuación ponderada de cada dimensión y la puntuación '
     'global. Se muestra el nivel de riesgo (BAJO / MODERADO / ALTO / CRÍTICO), el desglose gráfico '
     'por dimensión, los factores de mayor riesgo y el protocolo policial de actuación recomendado.'),
    ('Paso 6 — Informe final',
     'Se genera un informe estructurado con todos los datos del caso, las puntuaciones obtenidas, '
     'la valoración indicador por indicador y el protocolo de actuación. El informe puede imprimirse '
     'en formato PDF (Imprimir / Guardar PDF) o copiarse al portapapeles para su integración en el atestado.'),
]

for titulo, descripcion in pasos:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(titulo)
    r.font.size = Pt(10); r.font.bold = True; r.font.name = 'Calibri'
    r.font.color.rgb = AZUL_CLARO
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(4)
    p2.paragraph_format.left_indent  = Cm(0.5)
    r2 = p2.add_run(descripcion)
    r2.font.size = Pt(10); r2.font.name = 'Calibri'; r2.font.color.rgb = NEGRO

doc.add_paragraph()

# ── Pie de página ─────────────────────────────────────────────────────────────
add_par(doc,
    'MGRPIC v3.0 · Instrumento de uso exclusivo para unidades de Policía Judicial · '
    'Los resultados tienen carácter orientativo · No es una herramienta judicial',
    size=8, italic=True, color=GRIS_TEXT,
    align=WD_ALIGN_PARAGRAPH.CENTER, sb=12, sa=0)

# ── Guardar ───────────────────────────────────────────────────────────────────
out = '/home/user/MGRPIC/mgrpic/ANEXO_MGRPIC.docx'
doc.save(out)
import os
print(f"Generado: {out} ({os.path.getsize(out):,} bytes)")
